"""extractor.py - PPTX slide text/notes extractor and slide PNG rasterizer.

Functions:
- extract_slide_text(pptx_path): Extracts bullet points and speaker notes per slide.
- convert_pptx_to_images(pptx_path, output_dir): Converts PPTX to PDF via PowerPoint COM
  (with LibreOffice fallback), then rasterizes slides to sequential PNGs via pdf2image.
"""

import os
import platform
import shutil
import subprocess
from pathlib import Path
from typing import Any
from pptx import Presentation
from pdf2image import convert_from_path

from config import POPPLER_BIN_DIR


def extract_slide_text(pptx_path: str | Path) -> list[dict[str, Any]]:
    """Reads a .pptx file via python-pptx and extracts text, bullets, and speaker notes.

    Returns:
        List[dict]: A list of dicts with keys:
            - 'slide_number': int (1-based index)
            - 'bullet_points': list[str] (all bullet points/text paragraphs)
            - 'speaker_notes': str (speaker notes if present)
    """
    pptx_file = Path(pptx_path).resolve()
    if not pptx_file.exists():
        raise FileNotFoundError(f"Presentation file not found at: {pptx_file}")

    prs = Presentation(str(pptx_file))
    slides_data: list[dict[str, Any]] = []

    for index, slide in enumerate(prs.slides, start=1):
        bullet_points: list[str] = []

        # Iterate through all shapes to extract textual content
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line and line not in bullet_points:
                        bullet_points.append(line)
            elif shape.has_table:
                # Extract text from tables if present
                for row in shape.table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        bullet_points.append(" | ".join(row_texts))

        # Extract speaker notes
        speaker_notes = ""
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            if notes_tf:
                speaker_notes = notes_tf.text.strip()

        slides_data.append({
            "slide_number": index,
            "bullet_points": bullet_points,
            "speaker_notes": speaker_notes,
        })

    return slides_data


def _convert_pptx_to_pdf_powerpoint(pptx_abs: str, pdf_abs: str) -> bool:
    """Converts PPTX to PDF using Windows PowerPoint COM via comtypes."""
    try:
        import comtypes.client

        # Initialize COM in current thread
        comtypes.CoInitialize()
        try:
            powerpoint = comtypes.client.CreateObject("PowerPoint.Application")
            powerpoint.WindowState = 2  # ppWindowMinimized = 2

            # Open presentation
            pres = powerpoint.Presentations.Open(pptx_abs, WithWindow=False)
            ppSaveAsPDF = 32
            pres.SaveAs(pdf_abs, ppSaveAsPDF)
            pres.Close()
            powerpoint.Quit()
            return True
        finally:
            comtypes.CoUninitialize()
    except Exception as e:
        print(f"[-] PowerPoint COM conversion encountered an issue: {e}")
        return False


def _convert_pptx_to_pdf_libreoffice(pptx_abs: str, output_dir: str) -> bool:
    """Converts PPTX to PDF using headless LibreOffice."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        for candidate in [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]:
            if os.path.exists(candidate):
                soffice = candidate
                break

    if not soffice:
        return False

    cmd = [soffice, "--headless", "--convert-to", "pdf", pptx_abs, "--outdir", output_dir]
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, check=True)
        return res.returncode == 0
    except Exception as e:
        print(f"[-] LibreOffice conversion failed: {e}")
        return False


def convert_pptx_to_images(pptx_path: str | Path, output_dir: str | Path) -> list[str]:
    """Converts presentation to PDF using PowerPoint COM (or LibreOffice),

    then rasterizes slides to sequential PNG images (slide_{n}.png) via pdf2image.

    Returns:
        list[str]: Sequential list of paths to generated slide_{n}.png files.
    """
    pptx_file = Path(pptx_path).resolve()
    out_dir = Path(output_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    if not pptx_file.exists():
        raise FileNotFoundError(f"PPTX file not found: {pptx_file}")

    temp_pdf = out_dir / f"{pptx_file.stem}_temp.pdf"
    pptx_abs = str(pptx_file)
    pdf_abs = str(temp_pdf)

    # 1. Convert PPTX to PDF
    converted = False
    if platform.system() == "Windows":
        converted = _convert_pptx_to_pdf_powerpoint(pptx_abs, pdf_abs)

    if not converted or not temp_pdf.exists():
        print("[*] Attempting conversion fallback via LibreOffice headless...")
        converted = _convert_pptx_to_libreoffice = _convert_pptx_to_pdf_libreoffice(pptx_abs, str(out_dir))
        # LibreOffice names the output file <stem>.pdf in output_dir
        lo_pdf = out_dir / f"{pptx_file.stem}.pdf"
        if lo_pdf.exists():
            if temp_pdf.exists():
                temp_pdf.unlink()
            lo_pdf.rename(temp_pdf)
            converted = True

    if not converted or not temp_pdf.exists():
        raise RuntimeError(
            "Could not convert presentation to PDF. "
            "Ensure Microsoft PowerPoint is installed (on Windows) or LibreOffice is installed."
        )

    # 2. Rasterize PDF pages to PNG using pdf2image
    try:
        kwargs = {}
        if POPPLER_BIN_DIR:
            kwargs["poppler_path"] = POPPLER_BIN_DIR

        pages = convert_from_path(str(temp_pdf), dpi=150, **kwargs)
    except Exception as e:
        raise RuntimeError(
            f"Failed to rasterize PDF slides with pdf2image: {e}\n"
            f"Verify Poppler (pdftoppm) is installed and in PATH. Poppler dir: {POPPLER_BIN_DIR}"
        ) from e
    finally:
        # Clean up temporary PDF
        if temp_pdf.exists():
            try:
                temp_pdf.unlink()
            except OSError:
                pass

    slide_image_paths: list[str] = []
    for idx, page in enumerate(pages, start=1):
        img_path = out_dir / f"slide_{idx}.png"
        page.save(str(img_path), "PNG")
        slide_image_paths.append(str(img_path))

    return slide_image_paths
