"""create_demo_pptx.py - Helper script to generate a sample presentation for testing."""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN


def create_demo_presentation(filename: str = "demo_presentation.pptx") -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title Slide
    blank_layout = prs.slide_layouts[6]  # Blank
    slide1 = prs.slides.add_slide(blank_layout)

    tb1 = slide1.shapes.add_textbox(Inches(1), Inches(2), Inches(11.333), Inches(3.5))
    tf1 = tb1.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "The Future of AI Video Generation"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.alignment = PP_ALIGN.CENTER

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Autonomous pipelines turning slides into narrated presentations"
    p1_sub.font.size = Pt(24)
    p1_sub.alignment = PP_ALIGN.CENTER

    slide1.notes_slide.notes_text_frame.text = (
        "Welcome everyone! Today we are introducing a revolution in slide automation."
    )

    # Slide 2: Key Features
    slide2 = prs.slides.add_slide(blank_layout)
    tb2 = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(11.333), Inches(5.5))
    tf2 = tb2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = "Key Pipeline Architecture"
    p2.font.size = Pt(36)
    p2.font.bold = True

    bullets = [
        "1. PPTX Extraction: Bullet points and speaker notes via python-pptx",
        "2. High-Fidelity Rendering: PowerPoint COM to PDF rasterized by Poppler",
        "3. Intelligent Narration: Gemini 2.5 Flash structured script generation",
        "4. Edge-TTS Audio: Natural voice synthesis using en-US-GuyNeural",
        "5. MoviePy 2.0 Assembly: Sequential composition with smooth transitions",
    ]
    for b in bullets:
        p = tf2.add_paragraph()
        p.text = b
        p.font.size = Pt(20)

    slide2.notes_slide.notes_text_frame.text = (
        "Every single step in this pipeline is decoupled and completely modular."
    )

    # Slide 3: Conclusion & Next Steps
    slide3 = prs.slides.add_slide(blank_layout)
    tb3 = slide3.shapes.add_textbox(Inches(1), Inches(2), Inches(11.333), Inches(3.5))
    tf3 = tb3.text_frame
    p3 = tf3.paragraphs[0]
    p3.text = "Ready for Production"
    p3.font.size = Pt(40)
    p3.font.bold = True
    p3.alignment = PP_ALIGN.CENTER

    p3_sub = tf3.add_paragraph()
    p3_sub.text = "Fast, reliable, and completely automated video generation."
    p3_sub.font.size = Pt(22)
    p3_sub.alignment = PP_ALIGN.CENTER

    slide3.notes_slide.notes_text_frame.text = (
        "Thank you for watching! Configure your GEMINI_API_KEY to unlock advanced AI narration."
    )

    out_path = Path(__file__).parent / filename
    prs.save(str(out_path))
    print(f"[+] Demo presentation created at: {out_path.resolve()}")
    return out_path


if __name__ == "__main__":
    create_demo_presentation()
